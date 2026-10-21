#!/usr/bin/env python3
"""Pruebas del pre cierre del contador con datos ficticios (el repo es público).

Arma un Excel con la misma forma que el pre cierre (portada, balance tributario
y análisis de cuentas), lo pasa por el extractor y revisa lo que lee el motor.

    python3 tools/probar_contable.py
"""
import json
import subprocess
import sys
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "tools"))
import extraer  # noqa: E402

bien = mal = 0


def probar(nombre, condicion, detalle=""):
    global bien, mal
    if condicion:
        bien += 1
        print("  ✓", nombre)
    else:
        mal += 1
        print("  ✗", nombre, detalle)


# Montos inventados y no redondos, para no coincidir con cifras reales.
FONDOS = [(datetime(2026, 1, 12), "TRANS A PERSONA DE PRU", 81750),           # el banco cortó el nombre
          (datetime(2026, 2, 12), "PAGO:CAFETERIA FICTICIA", 54650),
          (datetime(2026, 2, 20), "DEVOLUCION FONDO", -11300)]
FONDOS_INICIAL = 9876                                                          # viene de otro año
HONORARIOS = [(datetime(2026, 1, 30), "PAGO HONORARIOS ENERO - COFUNDADOR EJEMPLO", 1013579),
              (datetime(2026, 2, 27), "PAGO HONORARIOS FEBRERO - COFUNDADOR EJEMPLO", 1013579),
              (datetime(2026, 2, 10), "TRANSFERENCIA A PERSONA DE PRUEBA", 441484)]
PROVEEDORES = [(datetime(2026, 2, 3), "TRANSFERENCIA A PROVEEDOR EJEMPLO", 823456),  # ya tiene factura en el SII
               (datetime(2026, 3, 9), "TRANSF. A OTRO PROVEEDOR", 312335)]
PRESTAMO = [(datetime(2026, 1, 5), "TRANSFERENCIA DE SOCIO UNO", 6123457),
            (datetime(2026, 2, 6), "TRANSFERENCIA A SOCIO UNO", 7777)]         # salió hacia el socio y se sumó como préstamo
suma = lambda movs: sum(m for _, _, m in movs)
CUENTAS = [
    ("1.1.1010.10.01", "CAJA..", 513000, 0),
    ("1.1.1010.10.03", "FONDOS POR RENDIR", suma(FONDOS) + FONDOS_INICIAL + 11300, 11300),
    ("1.1.1010.20.01", "BANCO EJEMPLO", 0, 8765432),                   # el debe se calcula abajo para que cuadre
    ("1.1.1010.20.07", "BANCO DIGITAL", 51234, 83456),                 # sobregiro: queda en el pasivo
    ("1.1.1040.10.01", "CLIENTES NACIONALES", 3141592, 1111111),
    ("1.1.1040.40.04", "ANTICIPO HONORARIOS", suma(HONORARIOS), 0),
    ("1.1.1040.50.02", "ANTICIPO A PROVEEDORES", suma(PROVEEDORES), 0),
    ("1.1.1080.10.01", "DRONES", 765432, 0),
    ("2.1.1070.20.01", "PROVEEDORES NACIONALES", 412345, 934567),
    ("2.1.1070.30.01", "DIVIDENDOS POR PAGAR", 0, 54321),              # es deuda, no patrimonio
    ("2.1.1070.40.02", "ANTICIPO DE CLIENTES", 0, 321987),
    ("2.1.1070.50.01", "PRESTAMO SOCIO UNO", 0, suma(PRESTAMO)),
    ("2.1.1070.50.02", "CUENTA CORRIENTE SOCIO DOS", 0, 76543),
    ("2.4.1000.10.01", "CAPITAL", 0, 1050500),
    ("2.4.1500.20.01", "RESULTADOS ACUMULADOS", 23456, 0),             # tiene código: no es la fila de resultado
    ("2.4.1500.50.01", "DIVIDENDOS PROVISORIOS", 1234321, 0),
    ("3.1.1010.10.01", "VENTAS", 104321, 5212345),
    ("3.5.1090.10.01", "CORRECCION MONETARIA", 0, 57),
    ("4.5.1030.10.01", "HONORARIOS", 612345, 0),
    ("4.5.1030.10.02", "ARRIENDOS", 912345, 0),
    ("4.5.1030.10.40", "REMUNERACION SOCIO", 200001, 0),               # sueldo del socio ya anotado
    ("4.5.3000.10.01", "DEPRECIACION", 87654, 0),                      # fuera del EBITDA
    ("4.5.2100.10.01", "CORRECCION MONETARIA", 31, 0),
]
_desc = -sum(d - h for _, _, d, h in CUENTAS)
CUENTAS[2] = ("1.1.1010.20.01", "BANCO EJEMPLO", _desc, 8765432)
VENTAS, GASTOS, REM = 5212345 - 104321, 612345 + 912345, 200001
RESULTADO = VENTAS + 57 - GASTOS - REM - 87654 - 31
SUELDO_MES = 123457
ANALISIS = {
    # el título de la celda C8 viene copiado de otra hoja, como pasa en los reales: manda el nombre de la hoja
    "ANTICIPO PROVEEDORES": ("ANTICIPO HONORARIOS", PROVEEDORES),
    "ANTICIPO HONORARIOS": ("ANTICIPO HONORARIOS", HONORARIOS),
    "FONDOS POR RENDIR": ("FONDOS POR RENDIR", FONDOS),
    "PRESTAMO SOCIO UNO": ("PRESTAMO SOCIO UNO", PRESTAMO),
    "CLIENTES NACIONALES": ("CLIENTES NACIONALES", [(datetime(2026, 3, 1), "11.111.111-1CLIENTE FICTICIO", 2030481)]),
}


def armar_excel(destino: Path, saldo_malo=False, formula=False, plan_raro=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "Portada"
    ws["B5"], ws["B8"], ws["C8"] = "EMPRESA FICTICIA SPA", "Periodo", datetime(2026, 3, 31)
    b = wb.create_sheet("Balance")
    b["B2"], b["B3"] = "Balance Tributario", "Periodo Desde 01/03/2026 al 31/03/2026"
    for c, t in zip("BCDEFGHIJK", ["CTA", "Nombre Cuenta", "Debe", "Haber", "Deudor", "Acreedor", "Activo", "Pasivo", "Perdida", "Ganancia"]):
        b[f"{c}9"] = t
    r = 10
    tot = {k: 0 for k in "HIJK"}
    for cod, nom, debe, haber in CUENTAS:
        if plan_raro and cod.startswith("4.5.1030.10.02"):
            cod = "5.1.0000.00.01"
        deu, acr = max(debe - haber, 0), max(haber - debe, 0)
        act = pas = per = gan = 0
        if cod[0] in "12":
            act, pas = deu, acr
        else:
            per, gan = deu, acr
        for c, v in zip("BCDEFGHIJK", [cod, nom, debe, haber, deu, acr, act, pas, per, gan]):
            b[f"{c}{r}"] = v
        for k, v in zip("HIJK", [act, pas, per, gan]):
            tot[k] += v
        r += 1
    if formula:
        b["D10"] = "=511111+1889"  # sin valor calculado: así queda un Excel que nunca se abrió en Excel
    b[f"C{r}"] = "TOTALES PERIODO"
    res = tot["K"] - tot["J"]
    b[f"C{r + 1}"], b[f"I{r + 1}"], b[f"J{r + 1}"] = "RESULTADO GANANCIA", res, res
    for hoja, (titulo, movs) in ANALISIS.items():
        a = wb.create_sheet(hoja)
        a["B2"], a["B8"], a["C8"] = "Análisis de Cuentas", "Nombre Cuenta", titulo
        a["B12"], a["C12"], a["D12"] = "Fecha", "Detalle", "Monto"
        fila = 13
        extra = 0
        if hoja == "FONDOS POR RENDIR":
            a[f"B{fila}"], a[f"D{fila}"] = "Saldo inicial", FONDOS_INICIAL
            fila += 1
            extra = FONDOS_INICIAL
        for f, d, m in movs:
            a[f"B{fila}"], a[f"C{fila}"], a[f"D{fila}"] = f, d, m
            fila += 1
        saldo = suma(movs) + extra + (1 if saldo_malo and hoja == "FONDOS POR RENDIR" else 0)
        a[f"B{fila + 1}"], a[f"D{fila + 1}"] = "Saldo al 46112", saldo
    wb.save(destino)


def no_extrae(ruta):
    try:
        extraer.extraer_contable(ruta)
    except SystemExit as e:
        return str(e)
    return ""


with tempfile.TemporaryDirectory() as tmp:
    tmp = Path(tmp)
    extraer.PRIVADO = tmp  # no tocar datos/privado
    print("Extractor")
    armar_excel(tmp / "precierre_prueba.xlsx")
    extraer.extraer_contable(tmp / "precierre_prueba.xlsx")
    C = json.loads((tmp / "contable.json").read_text(encoding="utf-8"))
    probar("el Excel ficticio cuadra (debe = haber)", sum(d - h for _, _, d, h in CUENTAS) == 0)
    probar("lee el período desde la portada", C["desde"] == "2026-01-01" and C["hasta"] == "2026-03-31", (C["desde"], C["hasta"]))
    probar("el resultado calza con el del balance", C["resultado"] == RESULTADO, (C["resultado"], RESULTADO))
    probar("lee todas las cuentas (también «Resultados acumulados», que tiene código)", len(C["cuentas"]) == len(CUENTAS), len(C["cuentas"]))
    hojas = {a["hoja"]: a for a in C["analisis"]}
    probar("lee anticipos, fondos y préstamos de socios, y nada más", set(hojas) == {"ANTICIPO PROVEEDORES", "ANTICIPO HONORARIOS", "FONDOS POR RENDIR", "PRESTAMO SOCIO UNO"}, set(hojas))
    probar("de los préstamos de socios guarda solo el saldo", hojas["PRESTAMO SOCIO UNO"]["movs"] == [] and hojas["PRESTAMO SOCIO UNO"]["saldo"] == suma(PRESTAMO))
    probar("avisa de transferencias hacia el socio sumadas como préstamo", any("hacia el socio" in a for a in C["avisos"]), C["avisos"])
    probar("no hay otros avisos", len(C["avisos"]) == 1, C["avisos"])
    probar("el nombre de la hoja manda sobre el título copiado", hojas["ANTICIPO PROVEEDORES"]["saldo"] == suma(PROVEEDORES))
    fon = hojas["FONDOS POR RENDIR"]["movs"]
    probar("el «saldo inicial» entra sin fecha y no corta la lectura", any(m["fecha"] is None and m["monto"] == FONDOS_INICIAL for m in fon) and len(fon) == 4, fon)
    probar("guarda montos negativos (devoluciones)", any(m["monto"] == -11300 for m in fon))
    armar_excel(tmp / "precierre_malo.xlsx", saldo_malo=True)
    extraer.extraer_contable(tmp / "precierre_malo.xlsx")
    M = json.loads((tmp / "contable.json").read_text(encoding="utf-8"))
    probar("avisa si los movimientos no suman el saldo", any("FONDOS POR RENDIR" in a for a in M["avisos"]), M["avisos"])
    armar_excel(tmp / "precierre_formula.xlsx", formula=True)
    probar("no acepta un Excel con fórmulas sin su valor", "sin su valor" in no_extrae(tmp / "precierre_formula.xlsx"))
    armar_excel(tmp / "precierre_plan.xlsx", plan_raro=True)
    probar("no acepta un plan de cuentas que no sigue 1-2-3-4", "plan de cuentas" in no_extrae(tmp / "precierre_plan.xlsx"))
    armar_excel(tmp / "precierre_prueba.xlsx")
    extraer.extraer_contable(tmp / "precierre_prueba.xlsx")

    print("Motor")
    js = r"""
    const M = require(process.argv[1]);
    const C = JSON.parse(require('fs').readFileSync(process.argv[2], 'utf8'));
    const fin = { compras: [{ prov: 'PROVEEDOR EJEMPLO SPA', fecha: '2026-01-20', monto: 999991 }, { prov: 'OTRO PROVEEDOR SPA', fecha: '2026-05-02', monto: 312335 }] };
    const r = M.leerContable({ contable: C, finanzas: fin, modelo: { sueldo_socio_mes: %d, cofundador_glosa: 'cofundador ejemplo' } }, '2026-04-10');
    const otro = M.leerContable({ contable: C, finanzas: fin, modelo: {} }, '2027-02-01');
    const sin = M.leerContable({ contable: null, modelo: {} });
    console.log(JSON.stringify({ r: { ...r, espera: r.espera.map(e => ({ id: e.id, monto: e.monto, n: e.n, devoluciones: e.devoluciones, facturado: e.facturado, arrastre: e.arrastre })) }, otro: otro.etiqueta, sin }));
    """ % SUELDO_MES
    salida = subprocess.run(["node", "-e", js, str(RAIZ / "app" / "motor.js"), str(tmp / "contable.json")], capture_output=True, text=True)
    if salida.returncode:
        print(salida.stderr)
        sys.exit(1)
    o = json.loads(salida.stdout)
    r = o["r"]
    e = {x["id"]: x for x in r["espera"]}
    probar("sin pre cierre, el motor no inventa nada", o["sin"] is None)
    probar("ventas y gastos sin lo no operacional (depreciación, corrección monetaria)", r["ventas"] == VENTAS and r["gastos"] == GASTOS, (r["ventas"], r["gastos"]))
    probar("separa el sueldo del socio que anotó el contador", r["remSocio"] == REM and r["resultadoOperacional"] == VENTAS - GASTOS - REM, (r["remSocio"], r["resultadoOperacional"]))
    probar("cuenta 3 meses de enero a marzo", r["nMeses"] == 3 and r["etiqueta"] == "enero a marzo", (r["nMeses"], r["etiqueta"]))
    probar("un pre cierre de otro año dice el año", o["otro"] == "enero a marzo de 2026", o["otro"])
    probar("lo que viene de otros años no es gasto de este", e["fondos"]["arrastre"] == FONDOS_INICIAL and e["fondos"]["monto"] == suma(FONDOS), e["fondos"])
    probar("una devolución no se cuenta como pago", e["fondos"]["n"] == 2 and e["fondos"]["devoluciones"] == 1, e["fondos"])
    probar("un anticipo a un proveedor con factura en el SII hasta el corte no es «sin documento»",
           e["proveedores"]["facturado"] == 823456 and e["proveedores"]["monto"] == 312335 and r["facturado"] == 823456, e["proveedores"])
    sin_doc = suma(HONORARIOS) + 312335 + suma(FONDOS)
    probar("lo pagado sin documento del año", r["sinDoc"] == sin_doc, (r["sinDoc"], sin_doc))
    probar("separa lo pagado al cofundador", r["cofundador"] == 2 * 1013579, r["cofundador"])
    nombres = {x["nombre"]: x for x in r["contrapartes"]}
    probar("junta el nombre que el banco cortó con el completo", "Persona de Prueba" in nombres and nombres["Persona de Prueba"]["monto"] == 441484 + 81750 and not any(n.startswith("Persona de Pru") and n != "Persona de Prueba" for n in nombres), list(nombres))
    probar("marca al cofundador en la lista", any(x["cofundador"] for x in r["contrapartes"]) and not nombres["Persona de Prueba"]["cofundador"])
    probar("préstamos de socios (también la cuenta corriente) y dividendos provisorios", r["socios"]["prestamosTotal"] == suma(PRESTAMO) + 76543 and r["socios"]["dividendos"] == 1234321, r["socios"])
    B = r["balance"]
    probar("los dividendos por pagar son deuda", B["pasivos"] == 934567 - 412345 + 54321 + 321987 + suma(PRESTAMO) + 76543, B["pasivos"])
    probar("el patrimonio es lo que tiene menos lo que debe", B["patrimonio"] == B["activos"] - B["pasivos"], B)
    probar("el sobregiro resta de los bancos", B["bancos"] == (_desc - 8765432) - (83456 - 51234), B["bancos"])
    probar("patrimonio = capital − resultados acumulados − dividendos + resultado", B["patrimonio"] == 1050500 - 23456 - 1234321 + C["resultado"], (B["patrimonio"], C["resultado"]))
    E = r["ebitda"]
    sueldo = max(3 * SUELDO_MES, REM)
    probar("el sueldo del socio es el mayor entre el del plan y el anotado, no la suma", E["sueldoSocio"] == sueldo, E["sueldoSocio"])
    probar("EBITDA = ventas − gastos − sin documento − sueldo del socio", E["ebitda"] == VENTAS - GASTOS - sin_doc - sueldo, E)
    probar("el cofundador va aparte y no se cuenta dos veces", E["gastos"] + E["cofundador"] == GASTOS + sin_doc, E)

print(f"\n{bien} pruebas bien, {mal} con error")
sys.exit(1 if mal else 0)
