#!/usr/bin/env python3
"""
AerIA · extractores de fuentes.

Cada fuente de Aéreo entra a AerIA por un adaptador que la deja como JSON
normalizado en datos/. El motor (tools/construir.py) solo lee datos/.

    python3 tools/extraer.py finanzas  <informe.html | finanzas.json> [url del artifact]
    python3 tools/extraer.py plan      <lectura-claude-docs.json | plan.xml> [url del doc]
    python3 tools/extraer.py cursos    <carpeta del repo aereo-web>
    python3 tools/extraer.py licitaciones <licitaciones.json del radar>
    python3 tools/extraer.py comercial  <carpeta del respaldo de la colección «cotizaciones»>
    python3 tools/extraer.py contable   <pre cierre del contador .xlsx>

Los cuatro focos:
    PASADO   finanzas      informe de cartolas, SII y Mercado Pago (artifact)
             contable      pre cierre del contador (Excel): balance y resultado del año
    PRESENTE cursos        fechas y cupos publicados en aereo.cl (repo aereo-web)
    COMERCIAL comercial    cotizaciones de la pestaña Comercial (se editan en la página;
                           el dashboard local se importa ahí con el extractor)
             licitaciones  radar de Mercado Público (repo licitaciones-aereo)
    FUTURO   plan          AÉREO 2027 · Plan de rentabilidad (Claude Docs)
    MARCA    la identidad visual sale de aereo-web/assets/aereo-ui.css
"""
import json
import re
import subprocess
import sys
import xml.etree.ElementTree as ET
from datetime import datetime, timezone
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
# Enlace opcional a la fuente original. NUNCA escribirlo en este archivo: el repo
# es público y esos enlaces abren los informes privados de Aéreo.
URL = sys.argv[3] if len(sys.argv) > 3 else None
PUBLICO = RAIZ / "datos" / "publico"
PRIVADO = RAIZ / "datos" / "privado"


def ahora():
    return datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")


def guardar(destino: Path, datos: dict):
    destino.parent.mkdir(parents=True, exist_ok=True)
    destino.write_text(json.dumps(datos, ensure_ascii=False, indent=1), encoding="utf-8")
    try:
        nombre = destino.relative_to(RAIZ)
    except ValueError:
        nombre = destino
    print(f"  → {nombre}  ({destino.stat().st_size // 1024} KB)")


# --------------------------------------------------------------------------
# PASADO · finanzas
# --------------------------------------------------------------------------
def extraer_finanzas(origen: Path):
    texto = origen.read_text(encoding="utf-8")
    if origen.suffix == ".json":
        d = json.loads(texto)
    else:
        m = re.search(r'<script id="data" type="application/json">(.*?)</script>', texto, re.S)
        if not m:
            sys.exit("No encontré el bloque <script id=\"data\"> en el informe financiero.")
        d = json.loads(m.group(1))

    cats = {k: {"label": v["label"], "grupo": v["grupo"]} for k, v in d["cats"].items()}
    movs = [
        {
            "id": m["id"], "cta": m["cta"], "fecha": m["fecha"], "mes": m["mes"],
            "monto": m["monto"], "desc": m["desc"], "contra": m.get("contra") or "",
            "cat": m["cat"], "conc": m.get("conc") or "", "nota": m.get("nota") or "",
        }
        for m in d["movs"]
    ]
    compras = [
        {
            "prov": c["prov"], "folio": c["folio"], "fecha": c["fecha"], "mes": c["mes"],
            "monto": c["monto"], "pago": c["pago"], "medio": c.get("medio") or "",
            "det": c.get("det") or "",
        }
        for c in d["compras"]
    ]
    sii = [
        {
            "folio": s["folio"], "fecha": s["fecha"], "mes": s["mes"], "tipo": s["tipo"],
            "cliente": s.get("cliente") or s.get("razon") or "", "rut": s.get("rut") or "",
            "monto": s["monto"], "estado": s["estado"], "det": s.get("det") or "",
            "periodo": s.get("periodo") or "",
        }
        for s in d["sii"]
    ]
    st = d["stats"]
    # El SII y el registro interno pueden llegar más allá del corte de las cartolas
    sii_hasta = (st.get("fact_post") or {}).get("sii_hasta") or max((s["fecha"] for s in sii), default="")
    post = (d.get("reg_q") or {}).get("post")
    salida = {
        "fuente": "Informe «Finanzas Aéreo ago–sep 2026» (artifact en claude.ai)",
        "url": URL,
        "generado_fuente": d["generado"],
        "extraido": ahora(),
        "periodo": {"desde": "2026-08-01", "hasta": "2026-09-30"},
        "sii_hasta": sii_hasta,
        "registro_post": post,
        "cats": cats,
        "serie": d["serie"],
        "cuadratura": [
            {k: c[k] for k in ("cta", "cuenta", "desde", "hasta", "saldo_inicial", "abonos", "cargos", "saldo_cartola")}
            for c in d["cuadratura"]
        ],
        "hallazgos": d["hallazgos"],
        "ventas_prod": d["ventas_prod"],
        "precios": d["precios"],
        "cobros": st["cob"],
        "mp": st["mp"],
        "facturacion": st["sii"],
        "compras_resumen": st["compras"],
        "iva_julio": st["iva_jul"],
        "caja_minima": st["min_san"],
        "dias_bajos": st["dias_bajos"],
        "conf_cat": st["conf_cat"],
        "sii": sii,
        "compras": compras,
        "movs": movs,
    }
    guardar(PRIVADO / "finanzas.json", salida)


# --------------------------------------------------------------------------
# FUTURO · plan 2027 (documento de Claude Docs)
# --------------------------------------------------------------------------
def _texto(nodo):
    return re.sub(r"\s+", " ", "".join(nodo.itertext())).strip()


def extraer_plan(origen: Path):
    crudo = origen.read_text(encoding="utf-8")
    if crudo.lstrip().startswith("{"):
        j = json.loads(crudo)
        xml = (j.get("data") or j.get("value") or {}).get("xml") or j.get("xml")
    else:
        xml = crudo
    raiz = ET.fromstring(xml)
    titulo, secciones, actual = "", [], None
    for n in raiz:
        if n.tag == "paragraph" and n.get("heading"):
            if n.get("heading") == "1":
                titulo = _texto(n)
                continue
            actual = {"titulo": _texto(n), "parrafos": [], "listas": [], "tablas": []}
            secciones.append(actual)
            continue
        if actual is None:
            continue
        if n.tag == "paragraph":
            t = _texto(n)
            if t:
                actual["parrafos"].append(t)
        elif n.tag == "list":
            actual["listas"].append([_texto(i) for i in n.iter("listItem") if _texto(i)])
        elif n.tag == "table":
            filas = [[_texto(c) for c in r.findall("cell")] for r in n.findall("row")]
            actual["tablas"].append(filas)
    salida = {
        "fuente": "AÉREO 2027 · Plan de rentabilidad (Claude Docs)",
        "url": URL,
        "extraido": ahora(),
        "titulo": titulo,
        "secciones": secciones,
    }
    guardar(PRIVADO / "plan2027.json", salida)


# --------------------------------------------------------------------------
# PRESENTE · cursos publicados en aereo.cl
# --------------------------------------------------------------------------
LEER_JS = r"""
const fs = require('fs'), vm = require('vm');
const ctx = { window: {} };
vm.createContext(ctx);
for (const f of process.argv.slice(1)) vm.runInContext(fs.readFileSync(f, 'utf8'), ctx);
process.stdout.write(JSON.stringify(ctx.window));
"""


def extraer_cursos(web: Path):
    archivos = [web / "datos" / "cursos-fechas.js", web / "datos" / "n23-fechas.js"]
    for a in archivos:
        if not a.exists():
            sys.exit(f"No encontré {a}. ¿Es la carpeta del repo aereo-web?")
    w = json.loads(subprocess.check_output(["node", "-e", LEER_JS, *map(str, archivos)]))
    max_n1, max_n23 = w.get("AEREO_CUPOS_MAX", 6), w.get("AEREO_N23_MAX", 6)
    n1 = [
        {"inicio": c["s"], "fin": c["e"], "libres": c["cupos"], "max": max_n1,
         "vendidos": max_n1 - c["cupos"], "extra": bool(c.get("extra"))}
        for c in w["AEREO_CURSOS"]
    ]
    n23 = [
        {"inicio": c["l"], "fin": c["m"], "libres": c["cupos"], "max": max_n23,
         "vendidos": max_n23 - c["cupos"]}
        for c in w["AEREO_N23"]
    ]
    commit = subprocess.run(["git", "-C", str(web), "log", "-1", "--format=%h %cI"],
                            capture_output=True, text=True).stdout.strip()
    salida = {
        "fuente": "aereo.cl · datos/cursos-fechas.js y datos/n23-fechas.js (repo aereo-web)",
        "commit": commit,
        "extraido": ahora(),
        "precios_web": {"nivel1": 149000, "nivel23": 189000, "pack": 299000},
        "meses_agotados_2026": w.get("AEREO_AGOTADOS_HISTORICOS", []),
        "nivel1": n1,
        "niveles23": n23,
    }
    guardar(PUBLICO / "cursos.json", salida)


# --------------------------------------------------------------------------
# PRESENTE · radar de licitaciones
# --------------------------------------------------------------------------
ENCAJES = [
    ("capacitacion", "Curso o capacitación",
     r"capacitaci|curso|formaci[oó]n|entrenamiento|instrucci[oó]n|credencial"),
    ("asesoria", "Asesoría AOC / DGAC", r"\baoc\b|asesor[ií]a|consultor[ií]a|regulariz"),
    ("equipo", "Compra de equipos o insumos",
     r"^adq|adquisici|compra|suministro|insumo|repuesto|accesorio|bater[ií]a|p[oó]liza|seguro|"
     r"arriendo|software|equipamiento|equipos|mantenci|reparaci"),
    ("servicio", "Servicio aéreo",
     r"servicio|fotogrametr|levantamiento|topogr|termogr|inspecci|ortofoto|ortomosaico|lidar|"
     r"fotograf|filmaci|audiovisual|cobertura|mapeo|monitoreo|plan regulador"),
]


def _encaje(item):
    """Clasifica por el NOMBRE del proceso; la descripción solo desempata si el nombre no dice nada."""
    for texto in (item.get("nombre", ""), item.get("descripcion", "")):
        texto = (texto or "").lower().strip()
        for clave, etiqueta, patron in ENCAJES:
            if re.search(patron, texto):
                return clave, etiqueta
    return "otro", "Otro"


def extraer_licitaciones(origen: Path):
    d = json.loads(origen.read_text(encoding="utf-8"))
    abiertas = []
    for i in d["items"]:
        if i.get("estado") != "abierta":
            continue
        clave, etiqueta = _encaje(i)
        abiertas.append({
            "codigo": i["codigo"], "fuente": i["fuente"], "nombre": i["nombre"],
            "organismo": i.get("organismo") or "", "region": (i.get("region") or "").strip(),
            "monto": i.get("monto") or 0, "publicada": i.get("fecha_publicacion"),
            "cierre": i.get("fecha_cierre"), "ofertas": i.get("ofertas"), "url": i.get("url"),
            "encaje": clave, "encaje_label": etiqueta,
        })
    abiertas.sort(key=lambda x: x["cierre"] or "")
    salida = {
        "fuente": "Radar de licitaciones Aéreo (repo licitaciones-aereo · docs/data/licitaciones.json)",
        "generado_fuente": d.get("generado"),
        "extraido": ahora(),
        "total_seguidas": len(d["items"]),
        "abiertas": abiertas,
    }
    guardar(PUBLICO / "licitaciones.json", salida)


# --------------------------------------------------------------------------
# COMERCIAL · respaldo de las cotizaciones que viven en la base de la página
# --------------------------------------------------------------------------
def extraer_comercial(origen: Path):
    """Arma datos/privado/comercial.json desde un respaldo de la base de la página.

    El respaldo se saca con la herramienta ArtifactData (action «list», colección
    «cotizaciones», out_dir=<carpeta>): deja un JSON por cotización en
    <carpeta>/cotizaciones/<id>.json. Las cotizaciones se editan en la página
    (pestaña Comercial); este archivo solo es el respaldo que va dentro del
    paquete y que la página muestra mientras la base no responde.
    """
    carpeta = origen / "cotizaciones" if (origen / "cotizaciones").is_dir() else origen
    if not carpeta.is_dir():
        sys.exit(f"{origen} no es una carpeta con los JSON de la colección «cotizaciones».")
    cotizaciones = []
    for f in sorted(carpeta.glob("*.json")):
        d = json.loads(f.read_text(encoding="utf-8"))
        d = d.get("data", d)  # algunos respaldos envuelven el documento
        d.setdefault("id", f.stem)
        cotizaciones.append(d)
    destino = PRIVADO / "comercial.json"
    previo = json.loads(destino.read_text(encoding="utf-8")) if destino.exists() else {}
    previo.update({"fuente": "Base de la página AerIA (colección cotizaciones)", "extraido": ahora(), "cotizaciones": cotizaciones})
    previo.setdefault("clientes", [])
    guardar(destino, previo)


# --------------------------------------------------------------------------
# PASADO · contable (pre cierre del contador)
# --------------------------------------------------------------------------
# Del análisis de cuentas se guarda el detalle solo de los pagos que esperan documento
# (anticipos de honorarios y de proveedores, fondos por rendir): es lo que AerIA muestra.
# Los préstamos de socios se leen para comprobarlos, pero de ellos se guarda solo el saldo.
ESPERA_DOC = re.compile(r"anticipos?\s+(a\s+|de\s+)?(honorarios|proveedores)|fondos?\s+(por|a)\s+rendir", re.I)
PRESTAMO_SOCIO = re.compile(r"(pr[eé]stamos?|cuenta\s+corriente|cta\.?\s*cte\.?).*(socio|accionista)|(socio|accionista).*(pr[eé]stamo|cuenta\s+corriente)", re.I)
HACIA = re.compile(r"^(transferencias?|transf\.?|trans|tranf)\s+a\s+", re.I)
# Lo que quedó por pagar al corte (facturas de proveedores, boletas de honorarios e impuestos): sirve para no
# contar dos veces lo que el contador ya anotó y el banco paga después. De las boletas y los impuestos se guarda
# solo la fecha y el monto; de las facturas, también el detalle (proveedor y folio), para cruzarlas con los anticipos.
POR_PAGAR = re.compile(r"(honorarios|impuestos?)\s+por\s+pagar|^proveedores|proveedores\s+(nacionales|por\s+pagar)", re.I)


def _texto_celda(v):
    return str(v).strip() if v is not None else ""


def _fecha_iso(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if hasattr(v, "isoformat"):
        return v.isoformat()[:10]
    return None


def _numero(v):
    return round(float(v)) if isinstance(v, (int, float)) and not isinstance(v, bool) else 0


def _pesos(n):
    return ("−" if n < 0 else "") + "$" + f"{abs(round(n)):,}".replace(",", ".")


def extraer_contable(origen: Path):
    """Arma datos/privado/contable.json desde el pre cierre del contador (Excel).

    El pre cierre trae el balance tributario acumulado del año hasta la fecha de
    la portada, el estado de resultados y un «análisis de cuentas» por hoja (cada
    movimiento con fecha, detalle y monto). Se leen los valores guardados en el
    Excel (no las fórmulas) y se comprueba que todo cuadre.
    """
    try:
        from openpyxl import load_workbook
    except ImportError:
        sys.exit("Falta openpyxl: pip install openpyxl")
    wb = load_workbook(origen, data_only=True)
    wb_formulas = load_workbook(origen, data_only=False)

    # La fecha de corte: la celda a la derecha de «Periodo» en la portada (o en la primera hoja que la tenga)
    hasta = None
    for ws in wb.worksheets:
        for fila in ws.iter_rows():
            for c in fila:
                if _texto_celda(c.value).lower() == "periodo":
                    v = ws.cell(row=c.row, column=c.column + 1).value
                    if _fecha_iso(v):
                        hasta = _fecha_iso(v)
                        break
            if hasta:
                break
        if hasta:
            break
    if not hasta:
        sys.exit("No encontré la fecha del período («Periodo») en el pre cierre.")

    # Balance tributario: la hoja con la fila de títulos CTA | Nombre Cuenta | Debe | Haber | …
    balance, titulos, fila_t = None, {}, None
    for ws in wb.worksheets:
        for fila in ws.iter_rows():
            nombres = {_texto_celda(c.value).lower(): c.column for c in fila if c.value is not None}
            if "cta" in nombres and "nombre cuenta" in nombres and "debe" in nombres:
                balance, titulos, fila_t = ws, nombres, fila[0].row
                break
        if balance:
            break
    if not balance:
        sys.exit("No encontré el balance tributario (fila con «CTA», «Nombre Cuenta» y «Debe»).")
    col = {k: titulos.get(k) for k in ("cta", "nombre cuenta", "debe", "haber", "deudor", "acreedor", "activo", "pasivo", "perdida", "ganancia")}
    if not col["perdida"]:
        col["perdida"] = titulos.get("pérdida")
    # Un Excel con fórmulas que nunca se abrió en Excel no trae sus valores: todo quedaría en $0 sin aviso
    sin_valor = 0
    bf = wb_formulas[balance.title]
    for r in range(fila_t + 1, balance.max_row + 1):
        for k in ("debe", "haber", "deudor", "acreedor", "activo", "pasivo", "perdida", "ganancia"):
            if col.get(k) and balance.cell(row=r, column=col[k]).value is None and str(bf.cell(row=r, column=col[k]).value or "").startswith("="):
                sin_valor += 1
    if sin_valor:
        sys.exit(f"El balance tiene {sin_valor} celdas con fórmula pero sin su valor calculado. Abre el Excel en Excel, guárdalo y vuelve a intentar.")
    en_texto = [f"{balance.cell(row=r, column=col['cta']).value} {k}" for r in range(fila_t + 1, balance.max_row + 1)
                for k in ("debe", "haber", "deudor", "acreedor", "activo", "pasivo", "perdida", "ganancia")
                if col.get(k) and isinstance(balance.cell(row=r, column=col[k]).value, str) and balance.cell(row=r, column=col[k]).value.strip()
                and not str(bf.cell(row=r, column=col[k]).value or "").startswith("=")]
    if en_texto:
        sys.exit("Hay montos guardados como texto en el balance (" + "; ".join(en_texto[:3]) + "). Conviértelos a número en Excel y vuelve a intentar.")
    cuentas, resultado_libro = [], None
    for r in range(fila_t + 1, balance.max_row + 1):
        codigo = _texto_celda(balance.cell(row=r, column=col["cta"]).value)
        nombre = _texto_celda(balance.cell(row=r, column=col["nombre cuenta"]).value)
        # «RESULTADO GANANCIA» o «RESULTADO PERDIDA» es una fila sin código (una cuenta «Resultados acumulados» sí tiene)
        if not codigo and nombre.upper().startswith("RESULTADO"):
            vals = [balance.cell(row=r, column=c).value for c in col.values() if c]
            monto = next((v for v in vals if isinstance(v, (int, float))), None)
            if monto is not None:
                perdida = "PERDIDA" in nombre.upper() or "PÉRDIDA" in nombre.upper()
                # con «PERDIDA» el monto viene positivo; con «GANANCIA» la fórmula ya trae su signo
                resultado_libro = -abs(round(monto)) if perdida else round(monto)
            continue
        if not codigo or nombre.upper().startswith("TOTAL"):
            continue
        val = lambda k: _numero(balance.cell(row=r, column=col[k]).value) if col.get(k) else 0
        cuentas.append({"codigo": codigo, "nombre": nombre.rstrip(". "), "debe": val("debe"), "haber": val("haber"),
                        "deudor": val("deudor"), "acreedor": val("acreedor"), "activo": val("activo"), "pasivo": val("pasivo"),
                        "perdida": val("perdida"), "ganancia": val("ganancia")})
    if not cuentas or not any(c["debe"] or c["haber"] for c in cuentas):
        sys.exit("El balance no trae montos (todo en $0). Revisa que el Excel tenga sus valores guardados.")
    # AerIA lee el plan de cuentas por su primer dígito: 1 activos, 2 pasivos y patrimonio, 3 ingresos, 4 gastos
    raras = [c["codigo"] + " " + c["nombre"] for c in cuentas
             if (c["codigo"][:1] in "12" and (c["perdida"] or c["ganancia"])) or (c["codigo"][:1] in "34" and (c["activo"] or c["pasivo"]))
             or c["codigo"][:1] not in "1234"]
    if raras or not any(c["codigo"][:1] == "3" and c["ganancia"] for c in cuentas):
        sys.exit("El plan de cuentas no sigue 1 activos, 2 pasivos, 3 ingresos y 4 gastos; AerIA no sabría leerlo. Ejemplos: " + "; ".join(raras[:4]))
    resultado = sum(c["ganancia"] for c in cuentas) - sum(c["perdida"] for c in cuentas)
    cuadre_inv = sum(c["activo"] for c in cuentas) - sum(c["pasivo"] for c in cuentas)
    avisos = []
    if resultado_libro is None:
        avisos.append("no encontré la fila de resultado del balance para comprobarlo")
    elif resultado_libro != resultado:
        avisos.append(f"el resultado del balance ({_pesos(resultado_libro)}) no calza con ganancias − pérdidas ({_pesos(resultado)})")
    if cuadre_inv != resultado:
        avisos.append(f"activos − pasivos ({_pesos(cuadre_inv)}) no calza con el resultado ({_pesos(resultado)})")

    # Análisis de cuentas: hojas con la fila Fecha | Detalle | Monto y un «Saldo al …» al final
    analisis = []
    for ws in wb.worksheets:
        es_espera, es_prestamo = bool(ESPERA_DOC.search(ws.title)), bool(PRESTAMO_SOCIO.search(ws.title))
        es_por_pagar = bool(POR_PAGAR.search(ws.title)) and not re.search(r"anticipo", ws.title, re.I)
        if not (es_espera or es_prestamo or es_por_pagar):
            continue
        if re.search(r"\(\d+\)\s*$", ws.title):
            avisos.append(f"dejé fuera la hoja «{ws.title}», que parece una copia")
            continue
        cab = None
        for fila in ws.iter_rows():
            textos = {_texto_celda(c.value).lower(): c.column for c in fila if c.value is not None}
            if "fecha" in textos and "detalle" in textos and "monto" in textos:
                cab = (fila[0].row, textos["fecha"], textos["detalle"], textos["monto"])
                break
        if not cab:
            avisos.append(f"la hoja «{ws.title}» no tiene la fila Fecha | Detalle | Monto")
            continue
        r0, cf, cd, cm = cab
        movs, saldo = [], None
        for r in range(r0 + 1, ws.max_row + 1):
            f = ws.cell(row=r, column=cf).value
            texto = _texto_celda(f).lower()
            monto = ws.cell(row=r, column=cm).value
            if texto.startswith("saldo") and (re.search(r"inicial|anterior", texto) or not movs):
                # lo que viene de años anteriores (un «saldo inicial», o un «saldo al…» antes de los movimientos):
                # entra en la suma, sin fecha, y se sigue leyendo
                if isinstance(monto, (int, float)):
                    movs.append({"fecha": None, "detalle": "Saldo inicial", "monto": _numero(monto)})
                continue
            if texto.startswith("saldo"):
                saldo = _numero(monto)
                break
            if _fecha_iso(f) and isinstance(monto, (int, float)):
                movs.append({"fecha": _fecha_iso(f), "detalle": _texto_celda(ws.cell(row=r, column=cd).value), "monto": _numero(monto)})
        suma = sum(m["monto"] for m in movs)
        if saldo is not None and saldo != suma:
            avisos.append(f"en «{ws.title.strip()}» los movimientos suman {_pesos(suma)} y el saldo dice {_pesos(saldo)}")
        if es_prestamo:
            hacia = [m for m in movs if m["monto"] > 0 and HACIA.search(m["detalle"])]
            if hacia:
                avisos.append(f"en «{ws.title.strip()}» hay {len(hacia)} transferencias hacia el socio por {_pesos(sum(m['monto'] for m in hacia))} sumadas como préstamo: revisarlas con el contador")
        # el nombre de la hoja manda: la celda del título a veces viene copiada de otra hoja
        tipo = "espera" if es_espera else "prestamo" if es_prestamo else "por_pagar"
        if tipo == "prestamo":
            guardados = []
        elif tipo == "por_pagar" and not re.search(r"proveedores", ws.title, re.I):
            guardados = [{"fecha": m["fecha"], "detalle": "", "monto": m["monto"]} for m in movs]
        else:
            guardados = movs
        analisis.append({"hoja": ws.title.strip(), "tipo": tipo, "saldo": saldo if saldo is not None else suma, "movs": guardados})

    for a in avisos:
        print("  ¡ojo!", a)
    año = hasta[:4]
    guardar(PRIVADO / "contable.json", {
        "fuente": f"Pre cierre del contador al {hasta[8:10]}-{hasta[5:7]}-{año} ({origen.name})",
        "archivo": origen.name,
        "extraido": ahora(),
        "desde": f"{año}-01-01",  # el balance tributario acumula desde el 1 de enero
        "hasta": hasta,
        "resultado": resultado,
        "cuentas": cuentas,
        "analisis": analisis,
        "avisos": avisos,
    })

ADAPTADORES = {
    "finanzas": extraer_finanzas,
    "plan": extraer_plan,
    "cursos": extraer_cursos,
    "licitaciones": extraer_licitaciones,
    "comercial": extraer_comercial,
    "contable": extraer_contable,
}

if __name__ == "__main__":
    if len(sys.argv) not in (3, 4) or sys.argv[1] not in ADAPTADORES:
        print(__doc__)
        sys.exit(1)
    print(f"AerIA · extrayendo {sys.argv[1]}")
    ADAPTADORES[sys.argv[1]](Path(sys.argv[2]).expanduser().resolve())
